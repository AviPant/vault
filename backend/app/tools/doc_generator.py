# DOCX/XLSX Industrial Document Exporter
import os
from datetime import datetime
from typing import Dict, Any, Optional, List

# pyrefly: ignore [missing-import]
from docx import Document
# pyrefly: ignore [missing-import]
from docx.shared import Pt, Inches, RGBColor
# pyrefly: ignore [missing-import]
from docx.enum.text import WD_ALIGN_PARAGRAPH

import pandas as pd


BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUTPUT_DIR = os.path.join(BASE_DIR, "data", "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def _timestamp_filename(prefix: str, ext: str) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_prefix = "".join(c if c.isalnum() or c in "-_ " else "_" for c in prefix)[:50]
    return f"{safe_prefix}_{ts}{ext}"


def create_approval_note(
    title: str,
    content: str,
    metadata: Optional[Dict[str, str]] = None
) -> str:
    """
    Generate a formal MRPL-style Approval Note (.docx).

    Args:
        title: Document title (e.g., "Pipe Replacement Approval - Unit 42")
        content: Body text (can include multiple paragraphs separated by newlines)
        metadata: Optional dict with keys like "prepared_by", "department",
                  "reference_no", "date"

    Returns:
        Absolute path to the generated .docx file.
    """
    metadata = metadata or {}
    doc = Document()

    # --- Header ---
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_para.add_run("MANGALORE REFINERY AND PETROCHEMICALS LIMITED")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("APPROVAL NOTE")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph()  # Spacer

    # --- Metadata Table ---
    ref_no = metadata.get("reference_no", "AN-" + datetime.now().strftime("%Y%m%d-%H%M"))
    meta_items = [
        ("Reference No.", ref_no),
        ("Date", metadata.get("date", datetime.now().strftime("%d-%m-%Y"))),
        ("Prepared By", metadata.get("prepared_by", "V.A.U.L.T. AI System")),
        ("Department", metadata.get("department", "Engineering")),
    ]

    table = doc.add_table(rows=len(meta_items), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (key, value) in enumerate(meta_items):
        table.rows[i].cells[0].text = key
        table.rows[i].cells[1].text = value

    doc.add_paragraph()  # Spacer

    # --- Title ---
    title_para = doc.add_paragraph()
    run = title_para.add_run(f"Subject: {title}")
    run.bold = True
    run.font.size = Pt(11)

    doc.add_paragraph()  # Spacer

    # --- Body Content ---
    for paragraph_text in content.split("\n"):
        text = paragraph_text.strip()
        if not text:
            continue
        # Detect section headers (lines ending with colon or all caps)
        if text.endswith(":") or (text.isupper() and len(text) < 80):
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(11)
        else:
            p = doc.add_paragraph(text)
            p.style = doc.styles["Normal"]

    # --- Footer ---
    doc.add_paragraph()
    doc.add_paragraph()
    sign_table = doc.add_table(rows=2, cols=3)
    headers = ["Prepared By", "Reviewed By", "Approved By"]
    for i, header in enumerate(headers):
        sign_table.rows[0].cells[i].text = header
        sign_table.rows[1].cells[i].text = "\n\n________________"

    # --- Save ---
    filename = _timestamp_filename(title, ".docx")
    filepath = os.path.join(OUTPUT_DIR, filename)
    doc.save(filepath)
    print(f"[DOC_GENERATOR] Created approval note: {filepath}")
    return filepath


def create_calculation_sheet(
    title: str,
    data: Dict[str, Any],
    notes: Optional[str] = None
) -> str:
    """
    Generate an industrial Calculation Sheet (.xlsx).

    Args:
        title: Sheet title (e.g., "Burst Pressure Calculation")
        data: Dict of parameter names to values, or a dict with
              "columns" and "rows" for tabular data.
        notes: Optional text notes to include in a second sheet.

    Returns:
        Absolute path to the generated .xlsx file.
    """
    filename = _timestamp_filename(title, ".xlsx")
    filepath = os.path.join(OUTPUT_DIR, filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        # Main calculation data
        if "columns" in data and "rows" in data:
            df = pd.DataFrame(data["rows"], columns=data["columns"])
        else:
            df = pd.DataFrame(
                list(data.items()),
                columns=["Parameter", "Value"]
            )

        df.to_excel(writer, sheet_name="Calculations", index=False)

        # Metadata sheet
        meta_df = pd.DataFrame([
            {"Field": "Title", "Value": title},
            {"Field": "Generated By", "Value": "V.A.U.L.T. AI System"},
            {"Field": "Date", "Value": datetime.now().strftime("%d-%m-%Y %H:%M")},
            {"Field": "Notes", "Value": notes or "N/A"},
        ])
        meta_df.to_excel(writer, sheet_name="Metadata", index=False)

    print(f"[DOC_GENERATOR] Created calculation sheet: {filepath}")
    return filepath
